import io
from zipfile import BadZipFile

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.errors import ParseFailure
from app.models import Block


class PptxParser:
    def parse(self, data: bytes, filename: str) -> tuple[str, list[Block]]:
        try:
            prs = Presentation(io.BytesIO(data))
        except (PackageNotFoundError, BadZipFile) as exc:
            raise ParseFailure(f"pptx: corrupted file: {exc}") from exc
        blocks: list[Block] = []
        title = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if shape == slide.shapes.title:
                    blocks.append(Block(type="heading", text=text, level=1))
                    if not title:
                        title = text
                else:
                    blocks.append(Block(type="paragraph", text=text))
        if not blocks:
            raise ParseFailure("pptx: no text content found")
        return title or filename.rsplit(".", 1)[0], blocks
