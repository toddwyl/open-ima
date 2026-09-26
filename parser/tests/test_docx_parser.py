import io

import pytest
from docx import Document

from app.errors import ParseFailure
from app.parsers.docx_parser import DocxParser


def test_docx_headings_and_paragraphs(sample_docx):
    title, blocks = DocxParser().parse(sample_docx, "memo.docx")
    assert title == "DOCX 主标题"
    types = [(b.type, b.level) for b in blocks]
    assert types == [
        ("heading", 1),
        ("paragraph", 0),
        ("heading", 2),
        ("paragraph", 0),
    ]


def test_docx_title_fallback(sample_docx):
    doc = Document()
    doc.add_paragraph("没有标题,只有段落。")
    buf = io.BytesIO()
    doc.save(buf)
    title, blocks = DocxParser().parse(buf.getvalue(), "plain.docx")
    assert title == "plain"
    assert blocks[0].text == "没有标题,只有段落。"


def test_docx_table_block():
    doc = Document()
    doc.add_paragraph("表格前段落。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    buf = io.BytesIO()
    doc.save(buf)
    title, blocks = DocxParser().parse(buf.getvalue(), "table.docx")
    assert title == "table"
    table_blocks = [b for b in blocks if b.type == "table"]
    assert len(table_blocks) == 1
    assert table_blocks[0].text == "A1 | B1\nA2 | B2"


def test_docx_empty_document_raises():
    doc = Document()
    for _ in range(3):
        doc.add_paragraph("")
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(ParseFailure, match="docx: no content"):
        DocxParser().parse(buf.getvalue(), "empty.docx")
