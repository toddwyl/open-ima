import pytest
from fpdf import FPDF


@pytest.fixture(scope="session")
def sample_pdf() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(0, 10, "First paragraph of the pdf document.", new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 10, "Second paragraph of the pdf document.", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


@pytest.fixture(scope="session")
def sample_txt() -> bytes:
    return b"First line of the text document.\nSecond line of the text document."


@pytest.fixture(scope="session")
def sample_md() -> bytes:
    return (
        b"# Sample Document\n\n"
        b"First paragraph of the markdown document.\n\n"
        b"Second paragraph of the markdown document.\n"
    )


@pytest.fixture(scope="session")
def sample_html() -> bytes:
    return (
        b"<html><head><title>Sample Document</title></head>"
        b"<body><h1>Sample Document</h1>"
        b"<p>First paragraph of the html document.</p>"
        b"<p>Second paragraph of the html document.</p>"
        b"</body></html>"
    )


@pytest.fixture(scope="session")
def sample_docx() -> bytes:
    import io
    from docx import Document

    doc = Document()
    doc.add_heading("DOCX 主标题", level=1)
    doc.add_paragraph("DOCX 第一段正文。")
    doc.add_heading("DOCX 小节", level=2)
    doc.add_paragraph("DOCX 第二段正文。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
